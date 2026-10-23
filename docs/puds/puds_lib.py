# -*- coding: utf-8 -*-
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.image.image import Image as DocxImage

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DOCS = os.path.join(ROOT, "docs", "puds")
DIA = os.path.join(DOCS, "diagramas")
ASSETS = os.path.join(DOCS, "assets")
OUT_DOCX = os.path.join(DOCS, "PUDS-UML_Clinica_Vida_Salud_Grupo-32.docx")

NAVY = RGBColor(0x0B, 0x3C, 0x5D)
BLUE = RGBColor(0x0A, 0x6E, 0xBD)
TEAL = RGBColor(0x1B, 0x99, 0x8B)
GRAY = RGBColor(0x5B, 0x6B, 0x7A)
LIGHT_GRAY = RGBColor(0x8A, 0x9A, 0xAA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x1F, 0x29, 0x33)

BODY_FONT = "Calibri"
HEAD_FONT = "Calibri"
MONO_FONT = "Consolas"


def _set_run_font(run, name=BODY_FONT, size=11, bold=False, italic=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), name)


def _shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _cell_margins(cell, top=60, bottom=60, left=90, right=90):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for tag, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)


def _table_borders(table, color="B8D4EA", size="4"):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tblPr.append(borders)


def _row_no_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    hdr = OxmlElement("w:tblHeader")
    hdr.set(qn("w:val"), "true")
    trPr.append(hdr)


def _para_border(p, color="0A6EBD", edge="bottom", size="8", space="2"):
    pPr = p._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.append(pBdr)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), size)
    el.set(qn("w:space"), space)
    el.set(qn("w:color"), color)
    pBdr.append(el)


def _para_shade(p, fill="EAF4FB"):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def _field(paragraph, instr, placeholder="1", size=9, color=GRAY, bold=False):
    runs = []
    r1 = paragraph.add_run()
    fc1 = OxmlElement("w:fldChar")
    fc1.set(qn("w:fldCharType"), "begin")
    r1._r.append(fc1)
    runs.append(r1)

    r2 = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    r2._r.append(it)
    runs.append(r2)

    r3 = paragraph.add_run()
    fc2 = OxmlElement("w:fldChar")
    fc2.set(qn("w:fldCharType"), "separate")
    r3._r.append(fc2)
    runs.append(r3)

    runs.append(paragraph.add_run(placeholder))

    r5 = paragraph.add_run()
    fc3 = OxmlElement("w:fldChar")
    fc3.set(qn("w:fldCharType"), "end")
    r5._r.append(fc3)
    runs.append(r5)
    for r in runs:
        _set_run_font(r, size=size, color=color, bold=bold)


class Builder:
    def __init__(self):
        self.doc = Document()
        self.fig_n = 0
        self.tab_n = 0
        self._setup_page()
        self._setup_styles()
        self._setup_header_footer()

    def _setup_page(self):
        s = self.doc.sections[0]
        s.page_width = Cm(21.0)
        s.page_height = Cm(29.7)
        s.left_margin = Cm(2.5)
        s.right_margin = Cm(2.5)
        s.top_margin = Cm(2.4)
        s.bottom_margin = Cm(2.2)
        s.header_distance = Cm(1.1)
        s.footer_distance = Cm(1.0)
        s.different_first_page_header_footer = True

    def _setup_styles(self):
        styles = self.doc.styles
        normal = styles["Normal"]
        normal.font.name = BODY_FONT
        normal.font.size = Pt(11)
        normal.font.color.rgb = BLACK
        normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        normal.paragraph_format.line_spacing = 1.15
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.space_before = Pt(0)
        rPr = normal.element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.append(rFonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rFonts.set(qn(attr), BODY_FONT)

        for name, size, color, before, after in (
            ("Heading 1", 19, NAVY, 0, 14),
            ("Heading 2", 14.5, BLUE, 14, 7),
            ("Heading 3", 12, NAVY, 10, 4),
        ):
            st = styles[name]
            st.font.name = HEAD_FONT
            st.font.size = Pt(size)
            st.font.bold = True
            st.font.color.rgb = color
            st.paragraph_format.space_before = Pt(before)
            st.paragraph_format.space_after = Pt(after)
            st.paragraph_format.keep_with_next = True
            st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if name == "Heading 1":
                st.paragraph_format.page_break_before = True
            rPr = st.element.get_or_add_rPr()
            rFonts = rPr.find(qn("w:rFonts"))
            if rFonts is None:
                rFonts = OxmlElement("w:rFonts")
                rPr.append(rFonts)
            for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rFonts.set(qn(attr), HEAD_FONT)

        cap = styles["Caption"]
        cap.font.name = BODY_FONT
        cap.font.size = Pt(8.5)
        cap.font.italic = True
        cap.font.color.rgb = GRAY
        cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_before = Pt(3)
        cap.paragraph_format.space_after = Pt(11)

        st = styles.add_style("CodeBlock", WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = styles["Normal"]
        st.font.name = MONO_FONT
        st.font.size = Pt(8.5)
        st.paragraph_format.space_before = Pt(4)
        st.paragraph_format.space_after = Pt(8)
        st.paragraph_format.line_spacing = 1.0
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

        for name in ("List Bullet", "List Number"):
            st = styles[name]
            st.font.name = BODY_FONT
            st.font.size = Pt(11)
            st.paragraph_format.space_after = Pt(3)
            st.paragraph_format.line_spacing = 1.1

    def _setup_header_footer(self):
        s = self.doc.sections[0]
        header = s.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.paragraph_format.tab_stops.add_tab_stop(Cm(8.0), WD_TAB_ALIGNMENT.CENTER)
        hp.paragraph_format.tab_stops.add_tab_stop(Cm(16.0), WD_TAB_ALIGNMENT.RIGHT)
        hp.paragraph_format.space_after = Pt(2)
        r = hp.add_run("Clínica Vida Salud · PUDS-UML")
        _set_run_font(r, size=8, color=GRAY)
        r = hp.add_run("\tIngeniería de Requisitos SC")
        _set_run_font(r, size=8, color=GRAY)
        r = hp.add_run("\tGrupo 32")
        _set_run_font(r, size=8, color=GRAY)
        _para_border(hp, color="B8D4EA", edge="bottom", size="6", space="2")

        footer = s.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.paragraph_format.tab_stops.add_tab_stop(Cm(16.0), WD_TAB_ALIGNMENT.RIGHT)
        fp.paragraph_format.space_before = Pt(2)
        r = fp.add_run("Proyecto de Examen Parcial II · Documento PUDS-UML")
        _set_run_font(r, size=8, color=GRAY)
        r = fp.add_run("\tPágina ")
        _set_run_font(r, size=8, color=GRAY)
        _field(fp, " PAGE ", "1", size=8, color=GRAY)
        r = fp.add_run(" de ")
        _set_run_font(r, size=8, color=GRAY)
        _field(fp, " NUMPAGES ", "1", size=8, color=GRAY)
        _para_border(fp, color="B8D4EA", edge="top", size="6", space="2")

        first_header = s.first_page_header
        first_header.is_linked_to_previous = False
        for p in first_header.paragraphs:
            p.text = ""
        first_footer = s.first_page_footer
        first_footer.is_linked_to_previous = False
        for p in first_footer.paragraphs:
            p.text = ""

        settings = self.doc.settings.element
        update = OxmlElement("w:updateFields")
        update.set(qn("w:val"), "true")
        settings.append(update)

    def front_title(self, text, size=17, color=NAVY, after=10):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(after)
        r = p.add_run(text)
        _set_run_font(r, name=HEAD_FONT, size=size, bold=True, color=color)
        return p

    def h1(self, text):
        p = self.doc.add_heading(text, level=1)
        for r in p.runs:
            _set_run_font(r, name=HEAD_FONT, size=19, bold=True, color=NAVY)
        return p

    def h2(self, text):
        p = self.doc.add_heading(text, level=2)
        for r in p.runs:
            _set_run_font(r, name=HEAD_FONT, size=14.5, bold=True, color=BLUE)
        return p

    def h3(self, text):
        p = self.doc.add_heading(text, level=3)
        for r in p.runs:
            _set_run_font(r, name=HEAD_FONT, size=12, bold=True, color=NAVY)
        return p

    def p(self, text="", bold_lead=None, size=11, color=BLACK, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
          after=6, before=0, italic=False, keep=False):
        para = self.doc.add_paragraph()
        para.alignment = align
        para.paragraph_format.space_after = Pt(after)
        para.paragraph_format.space_before = Pt(before)
        para.paragraph_format.keep_with_next = keep
        if bold_lead:
            r = para.add_run(bold_lead)
            _set_run_font(r, size=size, bold=True, color=NAVY)
        r = para.add_run(text)
        _set_run_font(r, size=size, color=color, italic=italic)
        return para

    def bullets(self, items, size=11):
        for item in items:
            lead = None
            text = item
            if isinstance(item, tuple):
                lead, text = item
            para = self.doc.add_paragraph(style="List Bullet")
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            para.paragraph_format.space_after = Pt(3)
            if lead:
                r = para.add_run(lead)
                _set_run_font(r, size=size, bold=True, color=NAVY)
            r = para.add_run(text)
            _set_run_font(r, size=size, color=BLACK)

    def numbered(self, items, size=11):
        for item in items:
            lead = None
            text = item
            if isinstance(item, tuple):
                lead, text = item
            para = self.doc.add_paragraph(style="List Number")
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            para.paragraph_format.space_after = Pt(3)
            if lead:
                r = para.add_run(lead)
                _set_run_font(r, size=size, bold=True, color=NAVY)
            r = para.add_run(text)
            _set_run_font(r, size=size, color=BLACK)

    def note(self, text, title="Nota"):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Cm(0.3)
        p.paragraph_format.right_indent = Cm(0.3)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(9)
        _para_shade(p, "EAF4FB")
        _para_border(p, color="0A6EBD", edge="left", size="18", space="6")
        r = p.add_run(title + ". ")
        _set_run_font(r, size=9.5, bold=True, color=BLUE)
        r = p.add_run(text)
        _set_run_font(r, size=9.5, color=NAVY)

    def code(self, text):
        p = self.doc.add_paragraph(style="CodeBlock")
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.right_indent = Cm(0.2)
        p.paragraph_format.keep_together = True
        _para_shade(p, "F3F6F9")
        _para_border(p, color="D5E3F0", edge="left", size="12", space="4")
        lines = text.strip("\n").split("\n")
        for i, line in enumerate(lines):
            if i:
                p.add_run().add_break()
            r = p.add_run(line if line else " ")
            _set_run_font(r, name=MONO_FONT, size=8.5, color=RGBColor(0x24, 0x33, 0x42))

    def caption(self, kind, text):
        p = self.doc.add_paragraph(style="Caption")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = False
        r = p.add_run(kind + " ")
        _set_run_font(r, size=8.5, bold=True, color=BLUE, italic=True)
        if kind == "Figura":
            self.fig_n += 1
            _field(p, " SEQ Figura \\* ARABIC ", str(self.fig_n), size=8.5, color=BLUE, bold=True)
        else:
            self.tab_n += 1
            _field(p, " SEQ Tabla \\* ARABIC ", str(self.tab_n), size=8.5, color=BLUE, bold=True)
        r = p.add_run(". " + text)
        _set_run_font(r, size=8.5, italic=True, color=GRAY)
        return p

    def figure(self, filename, width_cm, caption, break_before=False, max_height_cm=None):
        if break_before:
            self.doc.add_page_break()
        path = os.path.join(DIA, filename)
        if not os.path.exists(path):
            path = os.path.join(ASSETS, filename)
        if max_height_cm:
            info = DocxImage.from_file(path)
            aspect = info.height / info.width
            if width_cm * aspect > max_height_cm:
                width_cm = max_height_cm / aspect
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(path, width=Cm(width_cm))
        self.caption("Figura", caption)

    def table(self, headers, rows, widths, caption=None, font_size=8.5, first_col_bold=False,
              header_fill="0A6EBD", align_center_cols=None, row_fills=None):
        if caption:
            self.caption("Tabla", caption)
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        _table_borders(table)
        for i, w in enumerate(widths):
            table.columns[i].width = Cm(w)
        hdr = table.rows[0]
        _repeat_header(hdr)
        _row_no_split(hdr)
        for i, h in enumerate(headers):
            cell = hdr.cells[i]
            cell.width = Cm(widths[i])
            _shade(cell, header_fill)
            _cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(h))
            _set_run_font(r, size=font_size, bold=True, color=WHITE)
        for idx, row in enumerate(rows):
            tr = table.add_row()
            _row_no_split(tr)
            for i, val in enumerate(row):
                cell = tr.cells[i]
                cell.width = Cm(widths[i])
                _cell_margins(cell)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                if row_fills and idx < len(row_fills) and row_fills[idx]:
                    _shade(cell, row_fills[idx])
                elif idx % 2 == 1:
                    _shade(cell, "F4F8FC")
                cell.text = ""
                lines = str(val).split("\n")
                for j, line in enumerate(lines):
                    p = cell.paragraphs[0] if j == 0 else cell.add_paragraph()
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.0
                    if align_center_cols and i in align_center_cols:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    else:
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    r = p.add_run(line)
                    _set_run_font(r, size=font_size, bold=(first_col_bold and i == 0),
                                  color=NAVY if (first_col_bold and i == 0) else BLACK)
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(0)
        spacer.paragraph_format.space_before = Pt(0)
        spacer.paragraph_format.line_spacing = 0.6
        return table

    def kv_table(self, pairs, caption=None, label_w=4.4, value_w=11.6):
        self.table(["Campo", "Descripción"], pairs, [label_w, value_w], caption=caption,
                   font_size=8.8, first_col_bold=True)

    def toc(self):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        _field(p, ' TOC \\o "1-3" \\h \\z \\u ', "Actualice el campo para generar el índice.", size=10, color=BLACK)
        p2 = self.doc.add_paragraph()
        p2.paragraph_format.space_after = Pt(4)
        r = p2.add_run("Si el índice aparece vacío, seleccione el campo y presione F9 o use «Actualizar tabla» en Word.")
        _set_run_font(r, size=8.5, italic=True, color=LIGHT_GRAY)

    def list_of(self, kind):
        instr = f' TOC \\h \\z \\c "{kind}" '
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        _field(p, instr, f"Actualice el campo para generar el índice de {kind.lower()}s.", size=10, color=BLACK)

    def page_break(self):
        self.doc.add_page_break()

    def save(self, path):
        self.doc.save(path)


def cover(b):
    p = b.doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(10)
    p.add_run().add_picture(os.path.join(ASSETS, "portada_banner.png"), width=Cm(16.0))

    p = b.doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("PROYECTO DE EXAMEN PARCIAL II")
    _set_run_font(r, size=12, bold=True, color=BLUE)

    p = b.doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("SISTEMA WEB DE GESTIÓN DE CITAS MÉDICAS")
    _set_run_font(r, name=HEAD_FONT, size=25, bold=True, color=NAVY)

    p = b.doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("Clínica Vida Salud")
    _set_run_font(r, name=HEAD_FONT, size=20, bold=True, color=TEAL)

    p = b.doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("DOCUMENTO PUDS-UML")
    _set_run_font(r, size=14, bold=True, color=NAVY)

    p = b.doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(22)
    r = p.add_run("Perfil · Marco teórico · Modelo de negocio · Requisitos · Análisis · Diseño · Implementación · Pruebas")
    _set_run_font(r, size=9.5, italic=True, color=GRAY)

    table = b.doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    rows = [
        ("Grupo", "32"),
        ("Integrantes", "Matías Franco · Lucas Cortez Cáceres"),
        ("Docente", "MSc. Ing. Angélica Gazón Cuéllar"),
        ("Asignatura", "Ingeniería de Requisitos SC"),
        ("Fecha", "Octubre de 2026"),
        ("Versión", "1.0"),
        ("Repositorio", "github.com/MattFrs12343/Cl-nica-Vida-Salud-Grupo-32"),
    ]
    for i, (k, v) in enumerate(rows):
        tr = table.add_row()
        _row_no_split(tr)
        c0, c1 = tr.cells
        c0.width = Cm(4.4)
        c1.width = Cm(9.6)
        _cell_margins(c0, top=80, bottom=80)
        _cell_margins(c1, top=80, bottom=80)
        c0.text = ""
        c1.text = ""
        p0 = c0.paragraphs[0]
        p1 = c1.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r0 = p0.add_run(k.upper())
        _set_run_font(r0, size=8.5, bold=True, color=GRAY)
        r1 = p1.add_run(v)
        _set_run_font(r1, size=10, color=NAVY)
        if i < len(rows) - 1:
            _para_border(p0, color="D5E3F0", edge="bottom", size="4", space="4")
            _para_border(p1, color="D5E3F0", edge="bottom", size="4", space="4")

    p = b.doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(22)
    r = p.add_run("Documento elaborado para la asignatura Ingeniería de Requisitos SC · Defensa: lunes 09 de noviembre de 2026")
    _set_run_font(r, size=8.5, italic=True, color=LIGHT_GRAY)


def front_matter(b):
    b.page_break()
    b.front_title("INFORMACIÓN DEL DOCUMENTO", size=18, after=8)
    b.p("Este documento consolida los artefactos de las disciplinas de Captura de Requisitos, Análisis, Diseño, Implementación y Pruebas del Sistema Web de Gestión de Citas Médicas de la Clínica Vida Salud. Se elaboró con un enfoque iterativo e incremental basado en el Proceso Unificado y notación UML 2.5.", size=10.5)
    info = [
        ("Título", "Sistema Web de Gestión de Citas Médicas — Clínica Vida Salud"),
        ("Tipo de documento", "PUDS-UML (Perfil, Requisitos, Análisis, Diseño, Implementación y Pruebas)"),
        ("Versión", "1.0"),
        ("Estado", "Documento de diseño y avance académico"),
        ("Grupo", "32"),
        ("Integrantes", "Matías Franco · Lucas Cortez Cáceres"),
        ("Docente", "MSc. Ing. Angélica Gazón Cuéllar"),
        ("Asignatura", "Ingeniería de Requisitos SC"),
        ("Periodo", "Segundo Examen Parcial · 14/10/2026 al 09/11/2026"),
        ("Repositorio", "https://github.com/MattFrs12343/Cl-nica-Vida-Salud-Grupo-32"),
    ]
    b.kv_table(info, caption="Ficha técnica del documento")
    b.note("Las figuras, tablas y diagramas son de elaboración propia. El documento se apoya en los lineamientos de RUP/UML descritos en el material de la asignatura y en la especificación funcional del proyecto.", "Alcance de las fuentes")

    b.h2("Resumen ejecutivo")
    b.p("La Clínica Vida Salud necesita reemplazar la gestión manual de citas por un sistema web que centralice pacientes, médicos, especialidades, horarios, citas, fichas, historia clínica, pagos y reportes. El sistema se implementa con PHP 8.1+, PostgreSQL 16, HTML5, CSS3 y JavaScript, siguiendo una arquitectura MVC manual y desplegándose en hosting cPanel con base de datos gestionada en Neon o Railway.")
    b.p("La solución incorpora seis roles, cupos diferenciados 50/50 para web y kiosco, validación de disponibilidad en el servidor, prevención de citas superpuestas, reconsultas prioritarias, emisión de fichas con QR, envío de pedidos por WhatsApp mediante enlace profundo y auditoría de catálogos. El documento demuestra la trazabilidad entre requisitos, casos de uso, clases de análisis, diseño de datos, implementación y pruebas.")

    b.page_break()
    b.front_title("TABLA DE CONTENIDO", size=18, after=12)
    b.toc()

    b.page_break()
    b.front_title("ÍNDICE DE FIGURAS", size=18, after=10)
    b.list_of("Figura")
    b.front_title("ÍNDICE DE TABLAS", size=18, after=10)
    b.list_of("Tabla")


b = Builder()

